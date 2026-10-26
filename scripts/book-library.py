"""Explicit ingestion only: python scripts/book-library.py index | search QUERY | download"""
import sys,os,json,re,hashlib,sqlite3,shutil,urllib.request,urllib.parse,codecs
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]; BOOK=Path(os.environ.get('WHALE_BOOK_ROOT',ROOT/'book'))
INBOX=BOOK/'未索引'; DONE=BOOK/'已索引'; INDEX=BOOK/'索引'; DB=INDEX/'library.sqlite'
sys.path.insert(0,str(ROOT/'.data/book-tools'))
try:
 from opencc import OpenCC
 converter=OpenCC('t2s'); norm=lambda s:converter.convert(s).lower()
except ImportError:
 norm=lambda s:s.lower()
TOPICS={'儒家修身':['仁义','仁義','修身','君子','中庸','论语','論語','孟子','大学','大學'], '学习教育':['学习','學習','读书','讀書','教育','知之','学而','學而'], '人际交往':['朋友','友情','交友','礼仪','禮儀','沟通','溝通','对话'], '情绪慰藉':['安慰','难过','悲伤','悲傷','孤独','孤獨','忧愁','憂愁','快乐','快樂','喜怒'], '文学诗词':['诗经','詩經','诗','詩','词','詞','对联','對聯','文学','文學'], '兵法谋略':['军事','軍事','兵法','兵','将','將','谋略','謀略','鬼谷','练兵','練兵','战船'], '历史人物':['史','传','傳','志','历史','歷史','朝','人物'], '政治治国':['治国','治國','政','民','君','尚书','尚書','春秋'], '礼制伦理':['礼记','禮記','孝','礼','禮','伦理','倫理'], '易学哲思':['周易','乾','坤','阴阳','陰陽','卦','哲学']}
def tokens(s):
 s=norm(s); out=[]
 for w in re.findall(r'[\u3400-\u9fff]+|[a-z0-9]+',s):
  out.extend([w[i:i+2] for i in range(len(w)-1)] if re.match('[\u3400-\u9fff]',w) and len(w)>1 else [w])
 return ' '.join(out)
def decode(b):
 if b.startswith((b'\xff\xfe',b'\xfe\xff')): return b.decode('utf-16'),'utf-16'
 for enc in ['utf-8-sig','gb18030']:
  try:return b.decode(enc),enc
  except UnicodeDecodeError:pass
 # Microsoft's CP936 maps a standalone 0x80 to EURO; Python's GBK codec omits that extension.
 def cp936_error(error):
  if error.object[error.start:error.end]==b'\x80':return ('€',error.end)
  raise error
 codecs.register_error('whale-cp936-euro',cp936_error)
 try:return b.decode('gbk',errors='whale-cp936-euro'),'cp936-ms-euro'
 except UnicodeDecodeError:pass
 raise ValueError('无法无损解码，需人工确认编码')
def extract(p):
 if p.suffix.lower() in ['.txt','.md']:return decode(p.read_bytes())
 if p.suffix.lower()=='.pdf':
  from pypdf import PdfReader
  pages=[page.extract_text() or '' for page in PdfReader(p).pages]
  counts=[len(re.findall('[\u3400-\u9fff]',t)) for t in pages]
  # Repeated watermarks (e.g. three Chinese characters per page) are not a text layer.
  if sum(n>=20 for n in counts)<max(1,len(pages)*0.2):raise ValueError('扫描PDF或只有水印文字：有效正文页不足20%，需OCR及校对后再索引')
  text='\n'.join(f'\n【PDF第{i+1}页】\n'+t for i,t in enumerate(pages))
  return text,'pdf-text'
 if p.suffix.lower()=='.docx':
  from docx import Document
  d=Document(p);return '\n'.join([x.text for x in d.paragraphs]+[' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows]),'docx'
 raise ValueError('旧版DOC请先另存为DOCX或TXT；原文件保留未索引区')
TOPICS.update({'仁爱':['仁者','仁愛','仁爱','仁义','仁義'],'诚信':['诚信','誠信','信义','信義','有信'],'孝悌':['孝','悌','父母'],'忠诚':['忠','忠诚','忠誠'],'礼仪':['礼仪','禮儀','礼','禮'],'谦逊':['谦','謙','谦逊','謙遜'],'宽恕':['恕','宽容','寬容'],'自省':['三省','自省','反求诸己','反求諸己'],'勤学':['学而','學而','好学','好學','勤学','勤學'],'知行':['知行','躬行','实践','實踐'],'劝诫':['劝','勸','戒','诫','誡'],'交友':['朋友','友朋','交友','故人'],'离别':['离别','離別','送别','送別','送君'],'思乡':['故乡','故鄉','乡愁','鄉愁','思归','思歸'],'亲情':['父母','兄弟','姐妹','儿女','兒女'],'爱情':['相思','情人','佳人','窈窕'],'孤独':['孤独','孤獨','寂寞','独坐','獨坐'],'悲伤':['悲','哀','伤心','傷心','泪','淚'],'喜悦':['喜','乐','樂','欢','歡'],'山水':['山','水','江','湖'],'田园':['田园','田園','农','農','耕'],'春景':['春','桃花','柳'],'夏景':['夏','荷','莲','蓮'],'秋景':['秋','枫','楓','霜'],'冬景':['冬','雪','寒'],'月夜':['月','夜','星'],'咏物':['咏','詠','梅','竹','菊'],'怀古':['怀古','懷古','古迹','古跡','旧都','舊都'],'边塞':['边塞','邊塞','塞上','关山','關山'],'战争':['战争','戰爭','战','戰','征伐'],'将帅':['将军','將軍','帅','帥','统兵','統兵'],'训练':['练兵','練兵','操练','操練','训练','訓練'],'谋略':['谋','謀','计','計','权谋','權謀'],'军纪':['军法','軍法','军令','軍令','纪律','紀律'],'舟船':['舟','船','舰','艦'],'器械':['火器','铳','銃','炮','弓','弩'],'治政':['为政','為政','治国','治國','政事'],'民生':['民生','百姓','饥','饑','饥荒','饑荒'],'经济':['赋税','賦稅','财','財','食货','食貨'],'制度':['制度','法令','官制','典章'],'司法':['刑','狱','獄','讼','訟'],'地理':['地理','郡','县','縣','疆域'],'人物传记':['传曰','傳曰','列传','列傳','生平'],'史事':['元年','二年','纪年','紀年','史曰'],'阴阳':['阴阳','陰陽','乾','坤','卦'],'天命':['天命','天道','命运','命運'],'处世':['处世','處世','进退','進退','穷达','窮達']})
NORMAL_TOPICS={k:[norm(w) for w in words] for k,words in TOPICS.items()}
def tags(text):
 text=norm(text)
 return [k for k,words in NORMAL_TOPICS.items() if any(w in text for w in words)] or ['综合典籍']
def connect():
 c=sqlite3.connect(DB);c.executescript('CREATE TABLE IF NOT EXISTS books(id TEXT PRIMARY KEY,title TEXT,path TEXT,sha TEXT,tags TEXT,encoding TEXT,chunks INTEGER);CREATE TABLE IF NOT EXISTS chunks(id TEXT PRIMARY KEY,book TEXT,heading TEXT,start INTEGER,end INTEGER,text TEXT,tags TEXT);CREATE VIRTUAL TABLE IF NOT EXISTS book_fts USING fts5(id UNINDEXED,terms,detail=none);CREATE VIRTUAL TABLE IF NOT EXISTS chunk_fts USING fts5(id UNINDEXED,book UNINDEXED,terms,detail=none);');return c
def index(only=None,source=None):
 for d in [INBOX,DONE,INDEX]:d.mkdir(parents=True,exist_ok=True)
 source=Path(source).resolve() if source else INBOX.resolve()
 if not source.is_relative_to(INBOX.resolve()) or not source.is_dir():raise ValueError('索引范围必须是未索引区内的现有目录')
 # Only this explicitly invoked command stages loose books. Runtime search never writes/moves/indexes.
 for p in BOOK.iterdir():
  if source!=INBOX.resolve():continue
  if only and p.stem not in only:continue
  if p.is_file() and p.name!='BOOK_HANDOVER.md' and p.suffix.lower() in ['.txt','.md','.doc','.docx','.pdf']:
   target=INBOX/p.name
   if target.exists():raise ValueError('未索引区同名文件已存在：'+p.name)
   shutil.move(str(p),target)
 c=connect();report=[]
 for p in sorted(source.rglob('*')):
  if not p.is_file() or p.suffix.lower() not in ['.txt','.md','.doc','.docx','.pdf']:continue
  if only and p.stem not in only:continue
  try:
   raw=p.read_bytes();sha=hashlib.sha256(raw).hexdigest();bid=sha[:16]
   old=c.execute('select path from books where id=?',(bid,)).fetchone()
   if old:
    target=BOOK/old[0]
    if not target.exists() and target.resolve().is_relative_to(DONE.resolve()):target.parent.mkdir(parents=True,exist_ok=True);shutil.move(str(p),target)
    report.append({'file':p.name,'sourcePath':str(p.relative_to(BOOK)),'status':'already-indexed'});continue
   text,enc=extract(p);text=text.replace('\r\n','\n').replace('\r','\n')
   if len(text.strip())<50:raise ValueError('可提取正文过少')
   title=p.stem;category=tags(title)[0];target=DONE/category/p.name;target.parent.mkdir(parents=True,exist_ok=True)
   if target.exists():raise ValueError('已索引区同名文件已存在，不覆盖')
   rows=[];heading='正文';offset=0;buffer='';start=1;line_no=0
   def flush():
    nonlocal buffer,start
    if not buffer.strip():return
    cid=f'{bid}:{len(rows)+1:06d}';rows.append((cid,bid,heading,start,line_no,buffer.strip(),json.dumps(tags(title+' '+heading+' '+buffer),ensure_ascii=False)));buffer='';start=line_no+1
   for line_no,line in enumerate(text.splitlines(),1):
    if len(line)<60 and re.match(r'^\s*(?:第[一二三四五六七八九十百千零〇\d]+[卷篇章回节]|卷[一二三四五六七八九十百千零〇\d]+|[\u3400-\u9fff]{1,15}第[一二三四五六七八九十百\d]+)',line):flush();heading=line.strip()
    for part in [line[i:i+900] for i in range(0,len(line),900)] or ['']:
     if len(buffer)+len(part)>1000:flush();start=line_no
     buffer+=part+'\n'
   flush()
   with c:
    c.execute('insert into books values(?,?,?,?,?,?,?)',(bid,title,str(target.relative_to(BOOK)),sha,json.dumps(tags(title),ensure_ascii=False),enc,len(rows)))
    c.execute('insert into book_fts values(?,?)',(bid,tokens(title+' '+' '.join(tags(title)))))
    c.executemany('insert into chunks values(?,?,?,?,?,?,?)',rows)
    c.executemany('insert into chunk_fts values(?,?,?)',((r[0],bid,tokens(title+' '+r[2]+' '+r[5]+' '+r[6])) for r in rows))
   shutil.move(str(p),target);report.append({'file':p.name,'sourcePath':str(p.relative_to(BOOK)),'status':'indexed','chunks':len(rows),'encoding':enc,'category':category})
  except Exception as e:report.append({'file':p.name,'sourcePath':str(p.relative_to(BOOK)),'status':'unindexed','error':str(e)})
  print(json.dumps({'file':p.name,'status':report[-1]['status']},ensure_ascii=True),file=sys.stderr,flush=True)
 catalog=[dict(zip(['id','title','path','sha256','tags','encoding','chunks'],r)) for r in c.execute('select * from books')]
 (INDEX/'书目.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2),encoding='utf8');(INDEX/'本次索引报告.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8');c.close();print(json.dumps({'books':len(catalog),'chunks':sum(r['chunks'] for r in catalog),'report':report},ensure_ascii=False))
def search(query):
 if not DB.exists():return {'hits':[]}
 c=sqlite3.connect('file:'+DB.as_posix()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row
 ts=list(dict.fromkeys(tokens(query[:1500]).split()))[:40]
 if not ts:return {'hits':[]}
 expression=' OR '.join('"'+t+'"' for t in ts)
 requested=[];normalized_query=norm(query)
 # Stored editions may include a dynasty/credit filename suffix; readers ask by the work's short title.
 explicit_books=[r[0] for r in c.execute('select id,title from books') if norm(r[1]).split('（')[0].split('-')[0].strip() in normalized_query]
 enriched=bool(c.execute("select 1 from sqlite_master where name='tag_assignments'").fetchone())
 if enriched:
  vocabulary=c.execute("select distinct dimension,tag from tag_assignments where entity='article'").fetchall()
  # Short author names must not match inside ordinary phrases (e.g. 方来 in 自远方来).
  requested=[(d,t) for d,t in vocabulary if t in normalized_query and (d!='作者' or re.search(r'(?<![\u3400-\u9fff])'+re.escape(t),normalized_query) or any(prefix+t in normalized_query for prefix in ['作者','诗人','查找','查询','寻找','引用','关于']))]
 matched=tags(query);tagterms=tokens(' '.join(matched if matched!=['综合典籍'] else [])+' '+query)
 if not explicit_books and not requested and matched==['综合典籍'] and not re.search('书|書|典故|原文|引用|出自|出处|出處|什么意思|何意',query):c.close();return {'hits':[]}
 expr=' OR '.join('"'+t+'"' for t in dict.fromkeys(tagterms.split()))
 bookids=[r[0] for r in c.execute('select id from book_fts where book_fts match ? order by bm25(book_fts) limit 6',(expr,))] if expr else []
 if explicit_books:bookids=explicit_books[:6]
 # Select article candidates using dimension-qualified tags before reading fragments.
 articleids=[]
 if enriched:
  if requested:
   clauses=' OR '.join('(dimension=? AND tag=?)' for _ in requested)
   params=[v for pair in requested for v in pair]
   scope=' AND id IN (select id from articles where book IN ('+','.join('?' for _ in explicit_books)+'))' if explicit_books else ''
   articleids=[r[0] for r in c.execute("select id from tag_assignments where entity='article' AND ("+clauses+")"+scope+" group by id having count(distinct dimension || ':' || tag)=? order by id limit 40",params+explicit_books+[len(requested)])]
  if not articleids and bookids:
   articleids=[r[0] for r in c.execute('select id from article_fts where article_fts match ? AND book IN ('+','.join('?' for _ in bookids)+') order by bm25(article_fts) limit 12',[expression]+bookids)]
 args=[expression];where=''
 if bookids:where=' AND book IN ('+','.join('?' for _ in bookids)+')';args+=bookids
 if articleids:
  # Article candidates already encode explicit authors/tags, so do not discard them by book shortlist.
  where=' AND id IN (select chunk from article_chunks where article IN ('+','.join('?' for _ in articleids)+'))';args=[expression]+articleids
 ids=[r[0] for r in c.execute('select id from chunk_fts where chunk_fts match ?'+where+' order by bm25(chunk_fts) limit 3',args)]
 hits=[]
 for cid in ids:
  r=c.execute('select c.*,b.title,b.path from chunks c join books b on b.id=c.book where c.id=?',(cid,)).fetchone();hits.append({k:r[k] for k in ['id','title','path','heading','start','end','text','tags']})
 for hit in hits:
  if articleids:
   hit['articles']=[dict(r) for r in c.execute('select a.id,a.heading,a.tags from articles a join article_chunks ac on ac.article=a.id where ac.chunk=? AND a.id IN ('+','.join('?' for _ in articleids)+') limit 3',[hit['id']]+articleids)]
 c.close();return {'tags':matched,'selectedBooks':bookids,'selectedArticles':articleids,'hits':hits}
def download():
 raise RuntimeError('古籍来源已改为中华典藏与OpenELib；旧GitHub批量下载已停用。请先在指定站点核实具体书目/版本，将原文件和来源记录保存到未索引区，按用户要求再索引。详见book/BOOK_HANDOVER.md。')

def retag():
 import runpy
 globals()['retag_titles']=sys.argv[2:]
 runpy.run_path(str(ROOT/'scripts/book-enrich.py'))['enrich'](globals())

if __name__=='__main__':
 sys.stdout.reconfigure(encoding='utf8')
 if sys.argv[1]=='index':index(sys.argv[2:])
 elif sys.argv[1]=='index-dir':index(source=INBOX/sys.argv[2])
 elif sys.argv[1]=='retag':retag()
 elif sys.argv[1]=='download':download()
 elif sys.argv[1]=='search':print(json.dumps(search(sys.argv[2]),ensure_ascii=False))
 elif sys.argv[1]=='emotion-search':
  import runpy
  print(json.dumps(runpy.run_path(str(ROOT/'scripts/book-emotion.py'))['search'](globals(),json.loads(sys.argv[2])),ensure_ascii=False))
 elif sys.argv[1]=='get':
  c=sqlite3.connect('file:'+DB.as_posix()+'?mode=ro',uri=True);c.row_factory=sqlite3.Row
  table='articles' if ':article:' in sys.argv[2] else 'chunks'
  r=c.execute(f'select c.*,b.title,b.path from {table} c join books b on b.id=c.book where c.id=?',(sys.argv[2],)).fetchone();print(json.dumps(dict(r) if r else {},ensure_ascii=False));c.close()
